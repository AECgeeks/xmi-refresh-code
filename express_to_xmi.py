from collections import defaultdict
import csv
import glob
import json
import os
from pathlib import Path
from typing import Mapping, Optional

import ifcopenshell
import ifcopenshell.express
import toposort
import re

import psd
import xmi
import xmi_parse

w = ifcopenshell.ifcopenshell_wrapper

MODIFY_WORKBOOK = Path("data") / "4.4 properties.xlsx"
MODIFY_SHEET = "2 - For later update"


def _read_modify_rows(path: str | Path, sheet_name: str) -> list[list[str]]:
    """Read property-modify rows from either an .xlsx workbook or a .csv file.

    Returns rows from row 2 onwards (skipping the header) in the same shape
    as openpyxl.iter_rows(values_only=True) — i.e. a list of column-string lists.
    The column layout for both inputs matches the "2 - For later update" sheet:
      0 initial_property_name
      2 initial_pset
      3 initial_type
      5 new_property_name
      7 new_pset
      8 new_type
    The CSV equivalent (2_for_later_update.csv) has the same column order.
    """
    p = Path(path)
    if p.suffix.lower() == ".csv":
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                with p.open(encoding=encoding, newline="") as f:
                    reader = csv.reader(f)
                    header = next(reader, None) or []
                    rows = [list(row) for row in reader]
                    # 2_for_later_update.csv (2026-05-19 vintage) inserts a
                    # "Localised definition to use in revision commit" column at
                    # position 7, pushing new_pset to 8 and new_type to 9.
                    # Drop it so positional indices match the legacy xlsx layout
                    # (where new_pset is at 7, new_type at 8).
                    if any("localised definition" in (h or "").lower() for h in header):
                        idx = next(
                            i for i, h in enumerate(header)
                            if "localised definition" in (h or "").lower()
                        )
                        rows = [r[:idx] + r[idx + 1:] for r in rows]
                    return rows
            except UnicodeDecodeError:
                continue
        raise UnicodeDecodeError("utf-8/cp1252", b"", 0, 0, f"Could not decode {p}")

    from openpyxl import load_workbook
    workbook = load_workbook(p, read_only=True, data_only=True)
    try:
        return [
            [cell_text(value) for value in row]
            for row in workbook[sheet_name].iter_rows(min_row=2, values_only=True)
        ]
    finally:
        workbook.close()

def attribute_ref(a):
    if isinstance(a, w.named_type):
        return a.declared_type().name()
    elif isinstance(a, w.aggregation_type):
        return attribute_ref(a.type_of_element())
    elif isinstance(a, w.simple_type):
        return a.declared_type()
    else:
        breakpoint()

class package_resolver:
    def resolve(self, entity_name, xmi_model) -> tuple[str, ...]:
        ...

class old_xmi_based_package_resolver(package_resolver):
    def __init__(self, orig_xmi):
        self.package_paths = {n.name: n.package_path for n in orig_xmi.package_hierarchy}

    def resolve(self, entity_name, xmi_model) -> tuple[str, ...]:
        return self.package_paths[entity_name]


class input_based_package_resolver(package_resolver):
    def __init__(self):
        self.memory = {}

    def resolve(self, entity_name, xmi_model) -> tuple[str, ...]:
        if v := self.memory.get(entity_name):
            pass
        else:
            v = input(f"Package for {entity_name}: ")
            self.memory[entity_name] = v
        return xmi_model.name.upper(), v


class deprecated_entity_provider:
    def is_deprecated(self, entity_name) -> bool:
        ...


class old_xmi_based_deprecated_entity_provider(deprecated_entity_provider):
    def __init__(self, orig_xmi):
        self.deprecated_entities = orig_xmi.deprecated_entities

    def is_deprecated(self, entity_name) -> bool:
        return entity_name in self.deprecated_entities
    
class json_file_based_deprecated_entity_provider(deprecated_entity_provider):
    def __init__(self, fn):
        self.deprecated_entities = json.load(open(fn, encoding='utf-8'))

    def is_deprecated(self, entity_name) -> bool:
        return entity_name in self.deprecated_entities

class nothing_deprecated_entity_provider(deprecated_entity_provider):
    def __init__(self):
        pass

    def is_deprecated(self, entity_name) -> bool:
        return False

def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).strip()

def lookup_by_name(model: xmi.Model, name: str):
    element = model.by_name(name)
    if element is None:
        raise KeyError(name)
    return element

def lookup_attribute(cls: xmi.Class, name: str):
    attrs = {attribute.name: attribute for attribute in cls.attributes}
    return attrs[name]

class converter:
    M : xmi.Model
    models: xmi.ModelSet

    def __init__(self, package_resolver : package_resolver, schema_name, filepath : Optional[str | Path] = None):
        self.schema_name = schema_name
        self.filepath = Path(filepath) if filepath else None
        if self.filepath:
            self.M = xmi.Model.load(self.filepath)
        else:
            self.M = xmi.Model(schema_name)
        self.models = xmi.ModelSet.from_model(self.M)
        self.package_resolver = package_resolver

    def add_model(self, model: xmi.Model) -> None:
        self.models.add(model)

    def pkg(self, item_name : str):
        c = self.M
        for depth, pk in enumerate(self.package_resolver.resolve(item_name, self.M)):
            if pk.upper() == self.schema_name.upper():
                c = self.M
            elif el := self.M.by_name(pk):
                c = el
            else:
                if depth == 1:
                    imported_model = xmi.Model(pk)
                    self.add_model(imported_model)
                    c = imported_model.package(pk)
                    self.M.elements.append(xmi.Import(f"{pk}.uml#pk_{pk}", owner=self.M, imported=c))
                else:
                    c = c.package(pk)
        return c

    def convert_express_to_xmi(self, schema_source, deprecated_provider : deprecated_entity_provider, name_filter : Optional[set[str]] = None):
        s = ifcopenshell.express.parse(schema_source)
        self.schema = s
        
        if base := self.M.by_name("base"):
            pass
        else:
            base = self.M.package("base")
            for simple in ("binary", "boolean", "integer", "logical", "number", "real", "string"):
                base.datatype(simple)
        
            base.cls("PropertyTableValueBase")

        lines = list(open(schema_source).readlines())
        functions, rules = (defaultdict(list) for _ in range(2))
        for i, [name, *_] in filter(lambda t: t[1], [(i, re.findall(r'\bFUNCTION\s*(\w+)', ln)) for i, ln in enumerate(lines)]):
            while True:
                functions[name].append(lines[i])
                if re.findall(r'\bEND_FUNCTION\s*;', lines[i]):
                    break
                i += 1

        for i, [name, *_] in filter(lambda t: t[1], [(i, re.findall(r'\bRULE\s*(\w+)\s*FOR\b', ln)) for i, ln in enumerate(lines)]):
            while True:
                rules[name].append(lines[i])
                if re.findall(r'\bEND_RULE\s*;', lines[i]):
                    break
                i += 1

        unique_aggregates = {}
        derived = defaultdict(list)
        wheres = defaultdict(list)
        uniques = defaultdict(list)
        in_derive = False
        in_where = False
        in_unique = False
        entity = None
        for i, l in enumerate(lines):
            if ent := re.findall(r'\bENTITY\s*(\w+)', l):
                entity = ent[0]
            elif ent := re.findall(r'\bTYPE\s*(\w+)\s*=', l):
                entity = ent[0]
            elif re.findall(r'\bEND_ENTITY\s*;', l) or re.findall(r'\bEND_TYPE\s*;', l):
                entity = None
                in_derive, in_where, in_unique = False, False, False
            elif re.findall(r'\bDERIVE\b', l):
                in_derive, in_where, in_unique = True, False, False
            elif re.findall(r'\bUNIQUE\b', l) and not re.findall(r'\bOF\s+UNIQUE\b', l):
                in_derive, in_where, in_unique = False, False, True
            elif re.findall(r'\bWHERE\b', l):
                in_derive, in_where, in_unique = False, True, False
            elif entity is not None and in_derive:
                derived[entity].append(l)
            elif entity is not None and in_where:
                wheres[entity].append(l)
            elif entity is not None and in_unique:
                uniques[entity].append(l)
            elif entity is not None and re.findall(r'LIST \[\d+:(\d+|\?)\] OF UNIQUE', l):
                attr_name = l.split(':', 1)[0].strip()
                level = re.split(r'\s+', l[:l.index("OF UNIQUE")]).count("OF")
                unique_aggregates[(entity, attr_name)] = level

        def aggregation_suffix(a, for_inverse=False, attr: Optional[tuple] = None, nesting_depth=0):
            if isinstance(a, w.aggregation_type) or for_inverse:
                t = (("B","S",None) if for_inverse else "ABLS")[a.type_of_aggregation()]
                if t is None:
                    return ""
                if attr is not None and attr in unique_aggregates and unique_aggregates[attr] == nesting_depth:
                    t += "U"
                l = "?" if a.bound1() == -1 else a.bound1()
                u = "?" if a.bound2() == -1 else a.bound2()
                return f"_{t}[{l}:{u}]" + ("" if for_inverse else aggregation_suffix(a.type_of_element(), nesting_depth=nesting_depth+1, attr=attr))
            else:
                return ""

        for i, di in enumerate((functions, rules)):
            for k, lns in di.items():
                if name_filter and k not in name_filter:
                    continue
                self.pkg(k).constraint(
                    name=k,
                    language="EXPRESS_FUNCTION" if i == 0 else "EXPRESS_RULE",
                    body="".join(lns)
                )

        def get_deps(decl):
            if isinstance(decl, w.entity):
                # don't add attributes because we get into circular deps
                return [*((decl.supertype().name(),) if decl.supertype() else ())]
            elif isinstance(decl, w.type_declaration):
                ty = decl.declared_type()
                while isinstance(ty, w.aggregation_type):
                    ty = ty.type_of_element()
                if ty and not isinstance(ty, w.simple_type):
                    if isinstance(ty, w.named_type):
                        return [ty.declared_type().name()]
                else:
                    print(ty)
                    return []
            elif isinstance(decl, w.enumeration_type):
                return []
            elif isinstance(decl, w.select_type):
                return [x.name() for x in decl.select_list()]
            else:
                breakpoint()

        def break_none(decl, deps):
            assert None not in deps
            return deps

        deps : dict[str, list[str]] = {
            decl.name(): break_none(decl, get_deps(decl)) for decl in s.declarations.values()
            if name_filter is None or decl.name() in name_filter
        }
        ents = toposort.toposort_flatten(deps)

        # Keep only those not already existing
        ents = [e for e in ents if self.M.by_name(e) is None]
        ent_names = set(ents)

        def name_of(decl):
            while isinstance(decl, w.aggregation_type):
                decl = decl.type_of_element()
            if isinstance(decl, (w.type_declaration, w.entity)):
                return decl.name()
            elif isinstance(decl, w.simple_type):
                return decl.declared_type()
            elif isinstance(decl, w.named_type):
                return name_of(decl.declared_type())
            else:
                breakpoint()

        def emit(e):
            print(e.name())
            if isinstance(e, w.entity):
                # if e.name() == "IfcVibrationIsolator":
                #     breakpoint()
                sup = self.M.by_name(e.supertype().name()) if e.supertype() else None
                self.pkg(e.name()).cls(e.name(), sup, is_abstract=e.is_abstract())
            elif isinstance(e, w.enumeration_type):
                self.pkg(e.name()).enumeration(e.name(), e.enumeration_items())
            elif isinstance(e, w.type_declaration):
                self.pkg(e.name()).datatype(f"{e.name()}{aggregation_suffix(e.declared_type())}", self.M.by_name(name_of(e.declared_type())))
                if strlen := next(filter(None, [re.findall(r'STRING\s*\((\d+)\)\s*(FIXED)?', l) for l in lines if e.name() in l]), None):
                    length, fixed = strlen[0]
                    dt : xmi.DataType = self.M.by_name(e.name())
                    op = "=" if fixed else "<="
                    dt.constraints.append(xmi.Constraint('strlen', 'OCL', f'self.size() {op} {length}', owner=self.M.by_name(e.name())))
            elif isinstance(e, w.select_type):
                self.pkg(e.name()).union(e.name(), [self.M.by_name(x.name()) for x in e.select_list()])
            else:
                breakpoint()

        for e in map(s.declarations.__getitem__, ents):
            emit(e)

        processed_symmetric_inverses = set()

        for e in (s.declarations[x] for x in ents if isinstance(s.declarations[x], w.entity)):
            ecl : xmi.Class = self.M.by_name(e.name())
            if deprecated_provider.is_deprecated(e.name()):
                ecl.comment = "DEPRECATED"

            for a in e.attributes():
                a_ref = attribute_ref(a.type_of_attribute())
                target = self.M.by_name(a_ref)
                if target is None:
                    emit(s.declarations[a_ref])
                    target = self.M.by_name(a_ref)
                assert target is not None
                attr_name = f"{a.name()}{aggregation_suffix(a.type_of_attribute(), attr=(e.name(), a.name()))}"
                if isinstance(target, (xmi.Class, xmi.UnionType)):
                    orig_target_type = target_type = s.declarations[attribute_ref(a.type_of_attribute())]
                    inv_name = ""
                    if isinstance(target, xmi.Class):
                        # assert symmetric, union
                        for inv in target_type.all_inverse_attributes():
                            if inv.entity_reference().name() == e.name() and inv.attribute_reference().name() == a.name():
                                while target_type.supertype() and inv.name() in (i.name() for i in target_type.supertype().all_inverse_attributes()):
                                    target_type = target_type.supertype()
                                if orig_target_type is target_type:
                                    # assert symmetric, not sub-super-type distinction
                                    inv_name = f"INV_{inv.name()}{aggregation_suffix(inv, for_inverse=True)}"
                                    processed_symmetric_inverses.add((target_type.name(), inv.name()))
                    ecl.directed_association(
                        a.name(),
                        target,
                        inv_name,
                        attr_name,
                        target_lower=0 if a.optional() else 1,
                        target_upper=1
                    )
                else:
                    ecl.attr(
                        attr_name,
                        target,
                        lower=0 if a.optional() else 1,
                        upper=1
                    )
                print(' ', a.name(), a.type_of_attribute())
            if ptype_attr := next((a for a in e.attributes() if a.name() == "PredefinedType"), None):
                enum = ptype_attr.type_of_attribute().declared_type()
                for enum_lit in enum.enumeration_items():
                    nm = f"{enum.name()}.{enum_lit}"
                    if el := self.M.by_name(nm):
                        el.super.append(self.M.by_name(e.name()))
                    else:
                        self.pkg(e.name()).cls(nm, self.M.by_name(e.name()))

        for e in (s.declarations[x] for x in ents if isinstance(s.declarations[x], w.entity)):
            num_parent_inverses = len(e.supertype().all_inverse_attributes()) if e.supertype() else 0
            for inv in e.all_inverse_attributes()[num_parent_inverses:]:
                if (e.name(), inv.name()) not in processed_symmetric_inverses:
                    inv_name = f"INV_{inv.name()}{aggregation_suffix(inv, for_inverse=True)}"
                    opp : xmi.Class = self.M.by_name(inv.entity_reference().name())
                    opp.directed_association(e.name() + "_" + inv.name(), self.M.by_name(e.name()), inv_name, f"({inv.attribute_reference().name()})", owned_attribute=False)
                    print(' ', e.name(), inv)

        for ent, vs in derived.items():
            for nm, ty, stmt in re.findall(r"(\w+)\s*:\s*(.+?)\s*:=\s*(.+?)\s*;", "".join(vs), re.DOTALL):
                if ent not in ent_names:
                    continue
                print(ent, nm, ty, stmt)
                self.M.by_name(ent).attr(nm, self.M.by_name(ty.split(" ")[-1]), derived=" := ".join((ty, stmt)))

        for i, di in enumerate((wheres, uniques)):
            for ent, vs in di.items():
                if ent not in ent_names:
                    continue
                for nm, body in re.findall(r"(\w+)\s*:\s*(.+?)\s*;", "".join(vs), re.DOTALL):
                    if nm in ("CorrectPredefinedType", "CorrectTypeAssigned"):
                        # these are automatically generated
                        # except for some exceptions :)
                        if "EventTriggerType" not in body:
                            continue
                    self.M.by_name(ent).rule(nm, body, language="EXPRESS_WHERE" if i == 0 else "EXPRESS_UNIQUE")

    def write_properties(self, pset_pattern):
        if os.path.isdir(pset_pattern):
            pset_pattern = os.path.join(pset_pattern, '*.xml')
        for fn in glob.glob(pset_pattern):
            def format_type(ty) -> tuple[xmi.TypeRef, int, int | None, Mapping]:
                if isinstance(ty, psd.PropertyDef):
                    return format_type(ty.property_type)
                elif isinstance(ty, psd.QtoDef):
                    return self.M.by_name({
                        "Q_AREA": "IfcAreaMeasure",
                        "Q_COUNT": "IfcCountMeasure",
                        "Q_LENGTH": "IfcLengthMeasure",
                        "Q_NUMBER": "IfcQuantityNumber",
                        "Q_TIME": "IfcQuantityTime",
                        "Q_VOLUME": "IfcQuantityVolume",
                        "Q_WEIGHT": "IfcQuantityWeight",
                    }[ty.qto_type]), 1, 1, {}
                elif isinstance(ty, psd.PropertySingle):
                    return self.M.by_name(ty.data_type), 0, 1, {}
                elif isinstance(ty, psd.PropertyList):
                    return self.M.by_name(ty.data_type), 1, None, {}
                elif isinstance(ty, psd.PropertyBounded):
                    return self.M.by_name(ty.data_type), 0, 1, {"comment": "bounded"}
                elif isinstance(ty, psd.PropertyEnum):
                    if en := self.M.by_name(ty.enum_name):
                        pass
                    else:
                        en = self.pkg(ty.enum_name).enumeration(ty.enum_name, ty.enum_items)
                    return en, 1, None, {}
                elif isinstance(ty, psd.PropertyReference):
                    return self.M.by_name(ty.ref_type), 1, 1, {}
                elif isinstance(ty, psd.PropertyTable):
                    n = ty.defining + ty.defined
                    if el := self.M.by_name(n):
                        pass
                    else:
                        el = self.pkg(pdef.name).cls(n, self.M.by_name("PropertyTableValueBase"))
                        el.attr("DefiningValue", self.M.by_name(ty.defining), 1, None)
                        el.attr("DefinedValue", self.M.by_name(ty.defined), 1, None)
                    return el, 1, None, {}
                else:
                    raise TypeError()
                
            pdef = psd.parse(open(fn, encoding='utf-8').read())
            print(pdef.name)

            classes_cleaned = []
            for rel in pdef.applicable_classes:
                if '/' in rel:
                    rel, ptype = rel.split('/')
                else:
                    ptype = None
                cls : xmi.Class = self.M.by_name(rel)
                if cls is None:
                    raise ValueError(f"On {pdef.name} the applicability {rel} cannot be resolved to a valid type in UML")
                if 'IfcTypeObject' in cls.supertype_names():
                    cls = self.M.by_name(cls.name.removesuffix('Style').removesuffix('Type'))
                if ptype is not None:
                    attr = next(a for a in cls.attributes if a.name == 'PredefinedType')
                    cls = self.M.by_name(f"{attr.type.name}.{ptype}")
                if cls not in classes_cleaned:
                    classes_cleaned.append(cls)

            cls = self.pkg(pdef.name).cls(pdef.name, comment=pdef.template_type)
            for itm in pdef.items:
                *pos_args, kwargs = format_type(itm)
                cls.attr(itm.name, *pos_args, **kwargs)

            for ref in classes_cleaned:
                self.pkg(pdef.name).dependency(cls, ref)

    def modify_properties(self, workbook_path: str | Path, sheet_name: str = MODIFY_SHEET):
        rows = _read_modify_rows(workbook_path, sheet_name)

        pset_renames = []
        seen_pset_renames = set()
        for row in rows:
            old_pset = row[2]
            new_pset = row[7]
            if old_pset != new_pset and (old_pset, new_pset) not in seen_pset_renames:
                seen_pset_renames.add((old_pset, new_pset))
                pset_renames.append((old_pset, new_pset))

        for old_pset, new_pset in pset_renames:
            try:
                lookup_by_name(self.M, old_pset).name = new_pset
            except KeyError:
                continue

        for row in rows:
            property_name = row[0]
            old_type = row[3]
            new_pset = row[7]
            new_type = row[8]
            if old_type != new_type:
                try:
                    cls = lookup_by_name(self.M, new_pset)
                    lookup_attribute(cls, property_name).type = lookup_by_name(self.M, new_type)
                except KeyError:
                    continue

        for row in rows:
            old_property = row[0]
            new_property = row[5]
            new_pset = row[7]
            if old_property != new_property:
                try:
                    cls = lookup_by_name(self.M, new_pset)
                    lookup_attribute(cls, old_property).name = new_property
                except KeyError:
                    continue

    def save(self):
        self.models.save(self.filepath.parent if self.filepath is not None else None)

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', help='Input UML file')
    parser.add_argument('--express', help='Express schema file')
    parser.add_argument('--entity-names', help='Express schema entity names to filter on', nargs='+')
    parser.add_argument('--ea', help='Enterprise Architect .xmi file to resolve containing package and deprecated entities')
    parser.add_argument('--json', help='JSON file containing a single top-level array of deprecated entity names')
    parser.add_argument('--psd', nargs='+', action='append', help='One or more directories, glob patterns or single XML file paths for property set definitions')
    parser.add_argument('--modify', nargs='?', const=str(MODIFY_WORKBOOK), help='Apply property updates from a workbook (.xlsx) or CSV file (.csv)')
    parser.add_argument('--modify-sheet', default=MODIFY_SHEET, help='Workbook sheet to use with --modify (ignored for CSV input)')
    parser.add_argument('--output', help='Output UML file')
    args = parser.parse_args()

    if args.modify and not args.input:
        parser.error('--modify requires --input')

    pr, de = None, None
    if args.ea:
        orig_xmi = xmi_parse.parse_xmi(args.ea)
        pr = old_xmi_based_package_resolver(orig_xmi)
        de = old_xmi_based_deprecated_entity_provider(orig_xmi)
    elif args.json:
        de = json_file_based_deprecated_entity_provider(args.json)
    if pr is None:
        pr = input_based_package_resolver()
    if de is None:
        de = nothing_deprecated_entity_provider()
    
    if args.input:
        cv = converter(pr, 'ifc4x3_add2', args.input)
    else:
        cv = converter(pr, 'ifc4x3_add2')
    if args.express:
        nf = None
        if args.entity_names:
            nf = set(args.entity_names)
        ms = cv.convert_express_to_xmi(args.express, de, name_filter=nf)
    if args.psd:
        for pset_pattern in (pattern for group in args.psd for pattern in group):
            cv.write_properties(pset_pattern)
    if args.modify:
        cv.modify_properties(args.modify, args.modify_sheet)
    cv.save()
    os._exit(0)
