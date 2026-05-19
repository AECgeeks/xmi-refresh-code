#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import load_workbook

import psd

MAIN_SHEET = "1 - For initial upload"
PENUM_SHEET = "1 - PENUMs for initial upload"

# CSV-mode defaults — Yousheng's normalized catalog (see ifc4x4_properties/)
MAIN_CSV = "1_for_initial_upload.csv"
PENUM_CSV = "1_penum_initial_upload.csv"

INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*]')


@dataclass(frozen=True, slots=True)
class UploadRow:
    row_number: int
    property_name: str
    applicable_type: str
    set_name: str
    definition: str
    data_type: str


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def read_sheet_rows(workbook, sheet_name: str) -> list[list[str]]:
    try:
        worksheet = workbook[sheet_name]
    except KeyError as e:
        sheets = ", ".join(workbook.sheetnames)
        raise KeyError(f"Workbook has no sheet {sheet_name!r}; available sheets: {sheets}") from e

    return [
        [cell_text(value) for value in row]
        for row in worksheet.iter_rows(values_only=True)
    ]


def read_csv_rows(path: Path) -> list[list[str]]:
    # Excel CSV exports are typically cp1252 (Windows-Latin-1) with optional BOM.
    # Try utf-8-sig first (handles BOM and UTF-8 sources), fall back to cp1252.
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            with path.open(encoding=encoding, newline="") as f:
                return [list(row) for row in csv.reader(f)]
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("utf-8/cp1252", b"", 0, 0, f"Could not decode {path}")


def csv_main_to_excel_layout(rows: list[list[str]]) -> list[list[str]]:
    # 1_for_initial_upload.csv columns:
    #   0 property_name, 1 applicable_type, 2 pset, 3 definition, 4 data_type, 5 explanation
    # Excel "1 - For initial upload" columns (per load_upload_rows):
    #   0 property_name, 1 applicable_type, 2 pset, 3 definition, 4 definition_ext, 5 data_type
    # Inserts empty col 4 so data_type lands at col 5 where load_upload_rows expects it.
    if not rows:
        return []
    out = [rows[0]]
    for r in rows[1:]:
        padded = r + [""] * max(0, 5 - len(r))
        out.append([padded[0], padded[1], padded[2], padded[3], "", padded[4]])
    return out


def csv_penum_to_excel_layout(rows: list[list[str]]) -> list[list[str]]:
    # 1_penum_initial_upload.csv columns:
    #   0 index, 1 enum_name, 2 original_value, 3 initial_upload_value, 4 explanation, 5 original_desc
    # Excel "1 - PENUMs for initial upload" columns (per load_penums):
    #   0 _, 1 enum_name, 2 enum_value
    # Uses col 3 (initial_upload_value) — the canonical post-rename value.
    if not rows:
        return []
    out = [rows[0]]
    for r in rows[1:]:
        padded = r + [""] * max(0, 4 - len(r))
        out.append([padded[0], padded[1], padded[3]])
    return out


def normalize(value: str) -> str:
    return value.strip()


def load_upload_rows(rows: list[list[str]]) -> list[UploadRow]:
    upload_rows: list[UploadRow] = []
    for row_number, row in enumerate(rows[1:], start=2):
        padded = row + [""] * max(0, 6 - len(row))
        property_name = normalize(padded[0])
        applicable_type = normalize(padded[1])
        set_name = normalize(padded[2])
        definition = normalize(padded[3])
        definition_extension = normalize(padded[4])
        data_type = normalize(padded[5])

        if not any((property_name, applicable_type, set_name, definition, data_type)):
            continue

        missing = [
            name
            for name, value in (
                ("property name", property_name),
                ("applicable type", applicable_type),
                ("pset name", set_name),
                ("data type", data_type),
            )
            if not value
        ]
        if missing:
            raise ValueError(f"Row {row_number} is missing {', '.join(missing)}")

        upload_rows.append(
            UploadRow(
                row_number=row_number,
                property_name=property_name,
                applicable_type=applicable_type,
                set_name=set_name,
                definition=definition,
                data_type=data_type,
            )
        )

    return upload_rows


def load_penums(rows: list[list[str]]) -> dict[str, list[str]]:
    penums: dict[str, list[str]] = {}
    for row in rows[1:]:
        padded = row + [""] * max(0, 3 - len(row))
        enum_name = normalize(padded[1])
        enum_value = normalize(padded[2])
        if not enum_name or not enum_value:
            continue
        penums.setdefault(enum_name, [])
        if enum_value not in penums[enum_name]:
            penums[enum_name].append(enum_value)
    return penums


def split_applicable_types(value: str) -> list[str]:
    return [item.strip() for item in value.split(",") if item.strip()]


def unique_in_order(values) -> list[str]:
    result = []
    seen = set()
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result


def property_type_from_cell(row: UploadRow, penums: dict[str, list[str]]) -> psd.PropertyType:
    data_type = row.data_type
    if not data_type.casefold().startswith("penum_"):
        return psd.PropertySingle(data_type=data_type)

    penums_by_lower = {name.casefold(): values for name, values in penums.items()}
    enum_items = penums_by_lower.get(data_type.casefold())
    if enum_items is None:
        raise ValueError(f"Row {row.row_number} references {data_type}, but it is absent from the PENUM sheet")

    return psd.PropertyEnum(enum_name=data_type, enum_items=enum_items)


def build_definition_set(
    set_name: str,
    rows: list[UploadRow],
    penums: dict[str, list[str]],
    ifc_version: str,
    schema_location: str,
    qto_schema_location: str,
    include_definitions: bool,
) -> psd.DefinitionSet:
    applicable_classes = unique_in_order(
        item
        for row in rows
        for item in split_applicable_types(row.applicable_type)
    )

    if set_name.startswith("Pset_"):
        return psd.DefinitionSet(
            root_tag="PropertySetDef",
            template_type="PSET_TYPEDRIVENOVERRIDE",
            schema_location=schema_location,
            ifc_version=psd.IfcVersion(version=ifc_version),
            name=set_name,
            definition="",
            applicability="",
            applicable_classes=applicable_classes,
            applicable_type_value=",".join(applicable_classes),
            items_container_tag="PropertyDefs",
            item_tag="PropertyDef",
            items=[
                psd.PropertyDef(
                    name=row.property_name,
                    definition=row.definition if include_definitions else "",
                    property_type=property_type_from_cell(row, penums),
                )
                for row in rows
            ],
        )

    if set_name.startswith("Qto_"):
        return psd.DefinitionSet(
            root_tag="QtoSetDef",
            template_type="QTO_TYPEDRIVENOVERRIDE",
            schema_location=qto_schema_location,
            ifc_version=psd.IfcVersion(version=ifc_version),
            name=set_name,
            definition="",
            applicability="",
            applicable_classes=applicable_classes,
            applicable_type_value=",".join(applicable_classes),
            items_container_tag="QtoDefs",
            item_tag="QtoDef",
            items=[
                psd.QtoDef(
                    name=row.property_name,
                    definition=row.definition if include_definitions else "",
                    qto_type=row.data_type,
                )
                for row in rows
            ],
        )

    raise ValueError(f"{set_name} is not a Pset_* or Qto_* name")


def _load_inputs_from_xlsx(
    xlsx_path: Path,
    main_sheet: str,
    penum_sheet: str,
) -> tuple[list[UploadRow], dict[str, list[str]]]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        upload_rows = load_upload_rows(read_sheet_rows(workbook, main_sheet))
        penums = load_penums(read_sheet_rows(workbook, penum_sheet))
    finally:
        workbook.close()
    return upload_rows, penums


def _load_inputs_from_csv(
    main_csv: Path,
    penum_csv: Path,
) -> tuple[list[UploadRow], dict[str, list[str]]]:
    upload_rows = load_upload_rows(csv_main_to_excel_layout(read_csv_rows(main_csv)))
    penums = load_penums(csv_penum_to_excel_layout(read_csv_rows(penum_csv)))
    return upload_rows, penums


def convert(
    output_dir: Path,
    ifc_version: str,
    schema_location: str,
    qto_schema_location: str,
    include_definitions: bool,
    *,
    xlsx_path: Path | None = None,
    main_sheet: str = MAIN_SHEET,
    penum_sheet: str = PENUM_SHEET,
    main_csv: Path | None = None,
    penum_csv: Path | None = None,
) -> int:
    if xlsx_path is not None:
        upload_rows, penums = _load_inputs_from_xlsx(xlsx_path, main_sheet, penum_sheet)
    elif main_csv is not None and penum_csv is not None:
        upload_rows, penums = _load_inputs_from_csv(main_csv, penum_csv)
    else:
        raise ValueError("convert() requires either xlsx_path or both main_csv and penum_csv")

    groups: dict[str, list[UploadRow]] = {}
    for row in upload_rows:
        groups.setdefault(row.set_name, []).append(row)

    output_dir.mkdir(parents=True, exist_ok=True)
    for set_name, rows in groups.items():
        definition_set = build_definition_set(
            set_name=set_name,
            rows=rows,
            penums=penums,
            ifc_version=ifc_version,
            schema_location=schema_location,
            qto_schema_location=qto_schema_location,
            include_definitions=include_definitions,
        )
        psd.write(definition_set, output_dir / (set_name + ".xml"))

    return len(groups)


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert buildingSMART property workbook rows to PSD XML files.")
    parser.add_argument("output_dir", type=Path, help="Directory where generated XML files will be written")
    parser.add_argument("--xlsx", type=Path, default=None, help="Input .xlsx workbook (mutually exclusive with --csv-dir / --csv-main+--csv-penum)")
    parser.add_argument("--sheet", default=MAIN_SHEET, help="Workbook sheet containing property rows")
    parser.add_argument("--penum-sheet", default=PENUM_SHEET, help="Workbook sheet containing PEnum rows")
    parser.add_argument("--csv-dir", type=Path, default=None, help=f"Directory containing {MAIN_CSV} and {PENUM_CSV} (shortcut for --csv-main/--csv-penum)")
    parser.add_argument("--csv-main", type=Path, default=None, help="CSV file with property rows (initial-upload form)")
    parser.add_argument("--csv-penum", type=Path, default=None, help="CSV file with PEnum rows (initial-upload form)")
    parser.add_argument("--ifc-version", default=psd.DEFAULT_IFC_VERSION, help="IfcVersion value for generated PSD files")
    parser.add_argument("--schema-location", default=psd.DEFAULT_SCHEMA_LOCATION, help="PSD schema location")
    parser.add_argument("--qto-schema-location", default=psd.DEFAULT_QTO_SCHEMA_LOCATION, help="QTO schema location")
    parser.add_argument(
        "--include-definitions",
        action="store_true",
        help="Use column D as PropertyDef definitions; by default column D is ignored",
    )
    args = parser.parse_args()

    main_csv = args.csv_main
    penum_csv = args.csv_penum
    if args.csv_dir is not None:
        if main_csv is None:
            main_csv = args.csv_dir / MAIN_CSV
        if penum_csv is None:
            penum_csv = args.csv_dir / PENUM_CSV

    xlsx_given = args.xlsx is not None
    csv_given = main_csv is not None or penum_csv is not None
    if xlsx_given and csv_given:
        parser.error("Pass either --xlsx OR --csv-dir/--csv-main/--csv-penum, not both")
    if csv_given and (main_csv is None or penum_csv is None):
        parser.error("CSV mode requires both --csv-main and --csv-penum (or --csv-dir containing both)")
    if not xlsx_given and not csv_given:
        # Backwards-compatible default
        args.xlsx = Path("data") / "4.4 properties.xlsx"

    count = convert(
        output_dir=args.output_dir,
        xlsx_path=args.xlsx,
        main_sheet=args.sheet,
        penum_sheet=args.penum_sheet,
        main_csv=main_csv,
        penum_csv=penum_csv,
        ifc_version=args.ifc_version,
        schema_location=args.schema_location,
        qto_schema_location=args.qto_schema_location,
        include_definitions=args.include_definitions,
    )
    print(f"Wrote {count} PSD XML files to {args.output_dir}")


if __name__ == "__main__":
    main()
