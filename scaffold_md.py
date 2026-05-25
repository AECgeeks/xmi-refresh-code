"""MD scaffolding tool — generates skeleton .md docs from UML entity/enum names.

Usage:
    python scaffold_md.py <schema_uml_path> <docs_root> <entity_name> [<entity_name>...]

Example:
    python scaffold_md.py \\
        /path/to/schemas/IfcSharedBldgElements.uml \\
        /path/to/docs \\
        IfcArchElement IfcArchElementType IfcArchElementTypeEnum

Reads the UML to determine kind (uml:Class, uml:Enumeration,
uml:Constraint functions), attributes, and WHERE rules. Writes scaffolds to:
    {docs_root}/schemas/<layer>/<schema>/Entities/<Name>.md
    {docs_root}/schemas/<layer>/<schema>/Types/<Name>.md  (for enums)
    {docs_root}/schemas/<layer>/<schema>/Functions/<Name>.md  (for functions)
    {docs_root}/schemas/<layer>/<schema>/PropertySets/<Name>.md  (for psets)
    {docs_root}/schemas/<layer>/<schema>/QuantitySets/<Name>.md  (for qtos)
    {docs_root}/properties/<first-letter>/<Name>.md  (for pset/qto members)

Prose is intentionally placeholder — replace with real content from the
tunnel-team exec summary (or equivalent source) before merge.

Per `local/pipeline-changes.md` item E (Markdown scaffolding) and Thomas's
2026-05-01 ask about reducing the manual MD authoring burden.
"""
import argparse
import glob
import csv
import os
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from xml.dom import minidom

from xmi import Attribute, Class, Constraint, DataType, Enumeration, Model


# ---------------------------------------------------------------------------
# DocEntity.xml lookup — pull real attribute / WHERE-rule prose from the
# bSI-InfraRoom/IFC-Specification repo. Path layout:
#   <spec-dir>/IFC4x3/Sections/*/Schemas/*/Entities/<EntityName>/DocEntity.xml
# Entities new in IFC 4.4 (TM16, etc.) won't be present — caller falls back
# to convention-based synthesis (see _synth_attr_doc).
# ---------------------------------------------------------------------------


def _find_doc_entity_xml(spec_dir, entity_name):
    """Return the first DocEntity.xml path matching the entity, or None."""
    if not spec_dir:
        return None
    pattern = os.path.join(
        spec_dir, "IFC4x3", "Sections", "*", "Schemas", "*", "Entities",
        entity_name, "DocEntity.xml",
    )
    hits = glob.glob(pattern)
    return hits[0] if hits else None


def _text_of(node):
    """Concatenate all text-node children of an element."""
    if node is None:
        return ""
    parts = []
    for c in node.childNodes:
        if c.nodeType in (c.TEXT_NODE, c.CDATA_SECTION_NODE):
            parts.append(c.nodeValue or "")
    return "".join(parts)


def _clean_doc_text(s):
    """Normalise the raw <Documentation> text:
    - minidom already decodes &gt; / &lt; / &amp; / &apos; in text content.
    - But the DocEntity sources use a literal '&amp;nbsp;' sequence (an
      escaped HTML entity inside XML); after one round of XML decoding we
      see '&nbsp;' which we want as a real space.
    - Strip trailing whitespace per-line, drop a leading/trailing blank line.
    """
    if not s:
        return ""
    # &amp;nbsp; -> &nbsp; -> real space (NBSP would also be fine but breaks
    # plain-text diffs; collapse to a regular space).
    s = re.sub(r"(?:&nbsp;|\u00a0)+[ \t]*", " ", s)
    # Some sources double-escape: &amp;amp; → &amp; → keep as-is.
    # Normalise line endings, strip trailing whitespace.
    lines = [ln.rstrip() for ln in s.replace("\r\n", "\n").split("\n")]
    # Drop leading / trailing empties.
    while lines and not lines[0].strip():
        lines.pop(0)
    while lines and not lines[-1].strip():
        lines.pop()
    return "\n".join(lines)


def _read_text_file(path):
    if not path or not os.path.exists(path):
        return None
    try:
        return _clean_doc_text(Path(path).read_text(encoding="utf-8"))
    except Exception:
        return None


def _glob_first(*parts):
    hits = glob.glob(os.path.join(*parts))
    return hits[0] if hits else None


def _href_first_char(href):
    return href[0].lower() if href else ""


def _clean_xlsx_text(value):
    if value is None:
        return ""
    s = str(value)
    lines = [ln.rstrip() for ln in s.replace("\r\n", "\n").split("\n")]
    while lines and not lines[0].strip():
        lines.pop(0)
    while lines and not lines[-1].strip():
        lines.pop()
    return "\n".join(lines)


@dataclass(frozen=True)
class PropertyWorkbookRow:
    name: str
    applicable_type: str
    set_name: str
    definition: str
    local_definition: str
    data_type: str
    change_note: str


@dataclass(frozen=True)
class PEnumWorkbookItem:
    penum_name: str
    value: str
    description: str


@dataclass
class WorkbookContent:
    rows_by_set: dict
    penum_items: dict


def load_property_workbook(xlsx_path):
    """Read IFC 4.4 property/QTO/PEnum prose from the curated workbook."""
    empty = WorkbookContent(defaultdict(list), defaultdict(list))
    if not xlsx_path:
        return empty
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        return empty
    try:
        from openpyxl import load_workbook
    except Exception as e:
        print(
            f"[warn] {xlsx_path}: openpyxl is required to read property workbook ({e})",
            file=sys.stderr,
        )
        return empty

    try:
        wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    except Exception as e:
        print(f"[warn] failed to read {xlsx_path}: {e}", file=sys.stderr)
        return empty

    rows_by_set = defaultdict(list)
    if "1 - For initial upload" in wb.sheetnames:
        ws = wb["1 - For initial upload"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = _clean_xlsx_text(row[0] if len(row) > 0 else None)
            set_name = _clean_xlsx_text(row[2] if len(row) > 2 else None)
            if not name or not set_name:
                continue
            rows_by_set[set_name].append(
                PropertyWorkbookRow(
                    name=name,
                    applicable_type=_clean_xlsx_text(row[1] if len(row) > 1 else None),
                    set_name=set_name,
                    definition=_clean_xlsx_text(row[3] if len(row) > 3 else None),
                    local_definition=_clean_xlsx_text(row[4] if len(row) > 4 else None),
                    data_type=_clean_xlsx_text(row[5] if len(row) > 5 else None),
                    change_note=_clean_xlsx_text(row[6] if len(row) > 6 else None),
                )
            )

    penum_items = defaultdict(list)
    if "1 - PENUMs for initial upload" in wb.sheetnames:
        ws = wb["1 - PENUMs for initial upload"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            penum_name = _clean_xlsx_text(row[1] if len(row) > 1 else None)
            value = _clean_xlsx_text(row[2] if len(row) > 2 else None)
            if not penum_name or not value:
                continue
            penum_items[penum_name].append(
                PEnumWorkbookItem(
                    penum_name=penum_name,
                    value=value,
                    description=_clean_xlsx_text(row[3] if len(row) > 3 else None),
                )
            )

    return WorkbookContent(rows_by_set=rows_by_set, penum_items=penum_items)


def load_penum_norm_csv(csv_path):
    """{penum_name: [PEnumWorkbookItem]} from 2_penum_update_applied.csv.

    Uses post-normalization names ('Values for initial upload' column) so
    that [ED] scaffolds after apply_penum_normalization.py show correct names.
    """
    if not csv_path:
        return {}
    csv_path = Path(csv_path)
    if not csv_path.exists():
        return {}
    out = defaultdict(list)
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            name = (row.get("Penum Name") or "").strip()
            value = (row.get("Values for initial upload") or "").strip()
            desc = (row.get("Original Description") or "").strip()
            if name and value:
                out[name].append(PEnumWorkbookItem(name, value, desc))
    return dict(out)


def _extract_existing_short_def(md_path):
    """Return the text between the title and the short-definition marker."""
    if not md_path or not Path(md_path).exists():
        return None
    try:
        text = Path(md_path).read_text(encoding="utf-8")
    except Exception:
        return None
    marker = "<!-- end of short definition -->"
    if marker not in text:
        return None
    before_marker = text.split(marker, 1)[0]
    lines = before_marker.splitlines()
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    while lines and not lines[0].strip():
        lines.pop(0)
    while lines and not lines[-1].strip():
        lines.pop()
    short_def = "\n".join(lines).strip()
    return short_def or None


def load_short_def(spec_dir, name, kind):
    """Read Documentation.md from the canonical IFC-Specification location.

    kind in {entity, pset, qto, enum, penum, function}.
    PEnums embed their documentation inline in PropertyEnumerations/<x>/<name>.xml.
    """
    if not spec_dir:
        return None
    if kind == "penum":
        path = _glob_first(
            spec_dir, "IFC4x3", "PropertyEnumerations", "*", f"{name}.xml",
        )
        if not path:
            return None
        try:
            doc = minidom.parse(path)
        except Exception:
            return None
        for d in doc.getElementsByTagName("Documentation"):
            txt = _clean_doc_text(_text_of(d))
            if txt:
                return txt
        return None
    subdir = {
        "entity": "Entities",
        "pset": "PropertySets",
        "qto": "QuantitySets",
        "enum": "Types",
        "type": "Types",
        "function": "Functions",
    }.get(kind)
    if not subdir:
        return None
    path = _glob_first(
        spec_dir, "IFC4x3", "Sections", "*", "Schemas", "*", subdir,
        name, "Documentation.md",
    )
    return _read_text_file(path)


def load_enum_literal_docs(spec_dir, enum_name):
    """{literal_name: documentation} from DocEnumeration.xml → Constants/<x>/<href>.xml."""
    if not spec_dir:
        return {}
    path = _glob_first(
        spec_dir, "IFC4x3", "Sections", "*", "Schemas", "*", "Types",
        enum_name, "DocEnumeration.xml",
    )
    if not path:
        return {}
    try:
        doc = minidom.parse(path)
    except Exception:
        return {}
    out = {}
    for dc in doc.getElementsByTagName("DocConstant"):
        href = dc.getAttribute("href")
        if not href:
            continue
        const_path = os.path.join(
            spec_dir, "IFC4x3", "Constants", _href_first_char(href), f"{href}.xml",
        )
        if not os.path.exists(const_path):
            continue
        try:
            cdoc = minidom.parse(const_path)
        except Exception:
            continue
        for c in cdoc.getElementsByTagName("DocConstant"):
            nm = c.getAttribute("Name")
            doc_children = _children_by_tag(c, "Documentation")
            if nm and doc_children:
                txt = _clean_doc_text(_text_of(doc_children[0]))
                if txt:
                    out[nm] = txt
            break
    return out


def load_penum_literal_docs(spec_dir, penum_name):
    """{literal_name: documentation} from PropertyEnumerations → PropertyConstants/<x>/<href>.xml."""
    if not spec_dir:
        return {}
    path = _glob_first(
        spec_dir, "IFC4x3", "PropertyEnumerations", "*", f"{penum_name}.xml",
    )
    if not path:
        return {}
    try:
        doc = minidom.parse(path)
    except Exception:
        return {}
    out = {}
    for dc in doc.getElementsByTagName("DocPropertyConstant"):
        href = dc.getAttribute("href")
        if not href:
            continue
        const_path = os.path.join(
            spec_dir, "IFC4x3", "PropertyConstants", _href_first_char(href), f"{href}.xml",
        )
        if not os.path.exists(const_path):
            continue
        try:
            cdoc = minidom.parse(const_path)
        except Exception:
            continue
        for c in cdoc.getElementsByTagName("DocPropertyConstant"):
            nm = c.getAttribute("Name")
            doc_children = _children_by_tag(c, "Documentation")
            if nm and doc_children:
                txt = _clean_doc_text(_text_of(doc_children[0]))
                if txt:
                    out[nm] = txt
            break
    return out


def load_doc_entity(spec_dir, entity_name):
    """Return (attr_docs, rule_docs) dicts read from DocEntity.xml.

    attr_docs: {attribute_name: cleaned documentation string}
    rule_docs: {where_rule_name: cleaned documentation string}

    Returns ({}, {}) when the file is missing or unparseable.
    """
    path = _find_doc_entity_xml(spec_dir, entity_name)
    if not path:
        return {}, {}
    try:
        doc = minidom.parse(path)
    except Exception as e:
        print(f"[warn] failed to parse {path}: {e}", file=sys.stderr)
        return {}, {}

    attr_docs = {}
    for attrs_block in doc.getElementsByTagName("Attributes"):
        for da in _children_by_tag(attrs_block, "DocAttribute"):
            nm = da.getAttribute("Name")
            if not nm:
                continue
            doc_children = _children_by_tag(da, "Documentation")
            if not doc_children:
                continue
            txt = _clean_doc_text(_text_of(doc_children[0]))
            if txt:
                attr_docs[nm] = txt

    rule_docs = {}
    for rules_block in doc.getElementsByTagName("WhereRules"):
        for dr in _children_by_tag(rules_block, "DocWhereRule"):
            nm = dr.getAttribute("Name")
            if not nm:
                continue
            doc_children = _children_by_tag(dr, "Documentation")
            if not doc_children:
                continue
            txt = _clean_doc_text(_text_of(doc_children[0]))
            if txt:
                rule_docs[nm] = txt

    return attr_docs, rule_docs


# ---------------------------------------------------------------------------
# Convention-based synthesis for entities not yet in the IFC-Specification
# repo (new TM16/TM17/... entities). The upstream convention for
# `PredefinedType` is extremely consistent across IfcBeam / IfcColumn /
# IfcWall / IfcSlab / IfcMember / IfcPile etc., so we can synthesise the
# prose from the entity name alone.
# ---------------------------------------------------------------------------


def _entity_noun(entity_name):
    """IfcArchElement → 'arch element'; IfcBeam → 'beam'.

    Split CamelCase, drop the leading 'Ifc', lowercase. Multi-word entities
    get spaces (matching upstream phrasing like 'built element').
    """
    s = entity_name
    if s.startswith("Ifc"):
        s = s[3:]
    # CamelCase split
    s = re.sub(r"([a-z])([A-Z])", r"\1 \2", s)
    s = re.sub(r"([A-Z])([A-Z][a-z])", r"\1 \2", s)
    return s.lower()


def _indefinite_article(noun):
    """Best-effort a/an. Upstream uses 'a beam', 'a column'; vowel start → 'an'."""
    return "an" if noun[:1] in "aeiou" else "a"


def _synth_attr_doc(entity_name, attr_name, attr_type):
    """Synthesise upstream-convention prose for an attribute.

    Currently handles PredefinedType (the dominant case across TM16-shape
    entities). Returns None for attributes we don't have a convention for —
    caller falls back to the FILL-IN placeholder.
    """
    if attr_name == "PredefinedType":
        noun = _entity_noun(entity_name)
        article = _indefinite_article(noun)
        # Paired type entity name: IfcArchElement → IfcArchElementType.
        type_entity = f"{entity_name}Type"
        return (
            f"Predefined generic type for {article} {noun} that is specified in an "
            f"enumeration. There may be a property set given specifically for the "
            f"predefined types.\n"
            f"> NOTE  The _PredefinedType_ shall only be used, if no _{type_entity}_ "
            f"is assigned, providing its own _{type_entity}.PredefinedType_."
        )
    return None


def _synth_rule_doc(entity_name, rule_name):
    """Synthesise upstream-convention prose for a WHERE rule.

    Handles the two common rules paired with PredefinedType:
    CorrectPredefinedType and CorrectTypeAssigned.
    """
    noun = _entity_noun(entity_name)
    type_entity = f"{entity_name}Type"
    enum_name = f"{entity_name}TypeEnum"
    if rule_name == "CorrectPredefinedType":
        return (
            f"Either the _PredefinedType_ attribute is unset (e.g. because an "
            f"_{type_entity}_ is associated), or the inherited attribute "
            f"_ObjectType_ shall be provided, if the _PredefinedType_ is set to "
            f"USERDEFINED."
        )
    if rule_name == "CorrectTypeAssigned":
        return (
            f"Either there is no {noun} type object associated, i.e. the "
            f"_IsTypedBy_ inverse relationship is not provided, or the associated "
            f"type object has to be of type _{type_entity}_."
        )
    return None


def _children_by_tag(parent, tag):
    return [c for c in parent.childNodes if c.nodeType == c.ELEMENT_NODE and c.tagName == tag]


def _model_attrs(entity):
    """Return [(name, type_name)] for xmi.py Class attributes."""
    if not isinstance(entity, Class):
        return []
    return [
        (a.name, getattr(a.type, "name", "(unknown)"))
        for a in entity.attributes
        if isinstance(a, Attribute)
    ]


def _model_rules(entity):
    """Return [(name, body_first_line)] for xmi.py Class constraints."""
    if not isinstance(entity, Class):
        return []
    out = []
    for rule in entity.attributes:
        if not isinstance(rule, Constraint):
            continue
        preview = (rule.body or "").strip().splitlines()
        out.append((rule.name, preview[0][:140] if preview else ""))
    return out


def _model_enum_literals(entity):
    if not isinstance(entity, Enumeration):
        return []
    return [lit for lit in entity.literals if lit]


def scaffold_entity(name, supertype, attrs, rules, applicable_concepts=None,
                    attr_docs=None, rule_docs=None, short_def=None):
    """Build the .md scaffold body.

    short_def: optional Documentation.md text (canonical short definition).
    attr_docs / rule_docs: optional dicts of prose harvested from DocEntity.xml
    in the IFC-Specification repo. When a key is present, the real prose
    replaces the FILL-IN placeholder. When absent, we try convention-based
    synthesis (_synth_attr_doc / _synth_rule_doc) and only fall back to the
    FILL-IN placeholder if that also returns None.
    """
    attr_docs = attr_docs or {}
    rule_docs = rule_docs or {}
    lines = [f"# {name}", ""]
    if short_def:
        lines.append(short_def)
    else:
        lines.append(
            f"An _{name}_ represents <!-- FILL IN: short one-paragraph definition. "
            f"Source: tunnel-team exec summary or equivalent. -->"
        )
    lines.append("<!-- end of short definition -->")
    lines.append("")
    if attrs:
        lines.append("## Attributes")
        lines.append("")
        for a_name, a_type in attrs:
            lines.append(f"### {a_name}")
            if a_name in attr_docs:
                lines.append(attr_docs[a_name])
            else:
                synth = _synth_attr_doc(name, a_name, a_type)
                if synth is not None:
                    lines.append(synth)
                else:
                    lines.append(
                        f"<!-- FILL IN: description of {a_name} ({a_type}). -->"
                    )
            lines.append("")

    if rules:
        lines.append("## Formal Propositions")
        lines.append("")
        for r_name, r_body in rules:
            lines.append(f"### {r_name}")
            if r_name in rule_docs:
                lines.append(rule_docs[r_name])
            else:
                synth = _synth_rule_doc(name, r_name)
                if synth is not None:
                    lines.append(synth)
                else:
                    preview = (r_body or "<!-- FILL IN: rule description -->").replace("\n", " ")
                    lines.append(
                        f"<!-- FILL IN: prose explaining the rule. EXPRESS body: `{preview}` -->"
                    )
            lines.append("")

    # No '## Concepts' section emitted. For new entities (TM16+) without a
    # DocModelView.xml entry, the HTML builder auto-fills Concept content via
    # supertype-walk at render time. Explicit empty Concept headings would
    # override that auto-fill with nothing — worse than omitting the section.
    return "\n".join(lines) + "\n"


def scaffold_datatype(name, short_def=None):
    lines = [f"# {name}", ""]
    if short_def:
        first, _, rest = short_def.partition("\n\n")
        lines.append(first.strip())
        lines.append("<!-- end of short definition -->")
        if rest.strip():
            lines.append("")
            lines.append(rest.rstrip())
    else:
        lines.append(f"<!-- FILL IN: short definition of {name}. -->")
        lines.append("<!-- end of short definition -->")
    return "\n".join(lines) + "\n"


def scaffold_property_or_quantity_set(name, short_def=None, comments=None):
    comments = comments or []
    lines = [f"# {name}", ""]
    if short_def:
        lines.append(short_def)
    else:
        lines.append(f"<!-- FILL IN: short definition of {name}. -->")
    lines.append("<!-- end of short definition -->")
    lines.append("")

    if comments:
        lines.append("## Comments")
        lines.append("")
        for prop_name, text in comments:
            lines.append(f"### {prop_name}")
            lines.append("")
            lines.append(text)
            lines.append("")

    return "\n".join(lines).rstrip() + "\n"


def scaffold_normalized_property(name, definition):
    lines = [name, "=" * len(name), ""]
    if definition:
        lines.append(definition)
    else:
        lines.append(f"<!-- FILL IN: definition of {name}. -->")
    return "\n".join(lines).rstrip() + "\n"


def scaffold_enum(name, literals, short_def=None, literal_docs=None):
    literal_docs = literal_docs or {}
    lines = [f"# {name}", ""]
    if short_def:
        lines.append(short_def)
    else:
        lines.append(
            f"This enumeration defines the predefined types of <!-- FILL IN: which entity uses this enum -->."
        )
    lines.append("<!-- end of short definition -->")
    lines.append("")
    lines.append("## Items")
    lines.append("")
    for lit in literals:
        lines.append(f"### {lit}")
        if lit in literal_docs:
            lines.append(literal_docs[lit])
        else:
            lines.append(f"<!-- FILL IN: description of {lit}. -->")
        lines.append("")
    return "\n".join(lines) + "\n"


def scaffold_penum(name, items):
    lines = [f"# {name}", "<!-- end of short definition -->", "", "## Items", ""]
    for item in items:
        lines.append(f"### {item.value}")
        if item.description:
            lines.append(item.description)
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def _split_function_doc(markdown):
    """Split function Documentation.md into short definition and remainder."""
    markdown = (markdown or "").strip()
    if not markdown:
        return "", ""

    lines = markdown.splitlines()
    split_at = None
    saw_content = False
    for idx, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            if saw_content:
                split_at = idx
                break
            continue
        if saw_content and stripped.startswith(">"):
            split_at = idx
            break
        saw_content = True

    if split_at is None:
        return markdown, ""

    short_def = "\n".join(lines[:split_at]).strip()
    rest = "\n".join(lines[split_at:]).strip()
    return short_def, rest


def scaffold_function(name, short_def=None):
    doc_short, doc_rest = _split_function_doc(short_def)
    lines = [f"# {name}", ""]
    if doc_short:
        lines.append(doc_short)
    else:
        lines.append(f"<!-- FILL IN: short definition of {name}. -->")
    lines.append("<!-- end of short definition -->")
    if doc_rest:
        lines.append("")
        lines.append(doc_rest)
    return "\n".join(lines).rstrip() + "\n"


def _owning_schema_name(element):
    owner = getattr(element, "owner", None)
    while owner is not None:
        name = getattr(owner, "name", "")
        xmi_id = getattr(owner, "xmi_id", "")
        if name and xmi_id.startswith("pk_"):
            return name
        owner = getattr(owner, "owner", None)
    return None


def find_target_path(uml_path, docs_root, ent_name, kind, schema_name=None):
    """Best-effort path resolution: <docs_root>/schemas/<layer>/<schema>/<dir>/<name>.md.

    Existing docs are preferred when present. Otherwise <schema> is the UML
    filename stem and <layer> is read from <docs_root>/schemas/*/<schema>/.
    """
    docs_root = Path(docs_root)
    schema = schema_name or Path(uml_path).stem  # e.g. IfcSharedBldgElements
    subdir = {
        "entity": "Entities",
        "enum": "Types",
        "type": "Types",
        "function": "Functions",
        "penum": "PropertyEnumerations",
        "pset": "PropertySets",
        "qto": "QuantitySets",
    }[kind]

    existing = sorted((docs_root / "schemas").glob(f"*/*/{subdir}/{ent_name}.md"))
    if existing:
        return existing[0]

    # Fallback for legacy invocations that pass propertytypes.uml directly.
    if kind == "penum" and schema == "propertytypes":
        schema = "IfcSharedBldgElements"

    layer = None
    for cand in (docs_root / "schemas").iterdir() if (docs_root / "schemas").is_dir() else []:
        if (cand / schema).is_dir():
            layer = cand.name
            break
    if layer is None:
        layer = "shared"  # fallback

    return docs_root / "schemas" / layer / schema / subdir / f"{ent_name}.md"


def normalized_property_path(docs_root, name):
    first = name[:1].lower() if name else "_"
    return Path(docs_root) / "properties" / first / f"{name}.md"


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("uml_path", help="UML file containing the entity/enum")
    parser.add_argument("docs_root", help="Root docs directory (parent of schemas/)")
    parser.add_argument("names", nargs="+", help="Entity / enum names to scaffold")
    parser.add_argument("--force", action="store_true", help="Overwrite existing .md files")
    parser.add_argument("--dry-run", action="store_true", help="Don't write anything, just print")
    parser.add_argument(
        "--ifc-spec-dir",
        default=None,
        help=(
            "Path to bSI-InfraRoom/IFC-Specification checkout. When set, attribute "
            "and WHERE-rule prose is harvested from "
            "IFC4x3/Sections/*/Schemas/*/Entities/<EntityName>/DocEntity.xml. "
            "Entities not present (e.g. new TM16/TM17 entities) fall back to "
            "convention-based synthesis."
        ),
    )
    parser.add_argument(
        "--properties-xlsx",
        default=os.path.join("data", "4.4 properties.xlsx"),
        help=(
            "Workbook supplying IFC 4.4 property, quantity, and property-enumeration "
            "content. Uses sheets '1 - For initial upload' and "
            "'1 - PENUMs for initial upload'."
        ),
    )
    parser.add_argument(
        "--penum-norm-csv",
        default=None,
        help=(
            "Path to 2_penum_update_applied.csv. When set, PEnum MDs are built from "
            "post-normalization literal names in 'Values for initial upload' column "
            "instead of the workbook. Use this flag when re-scaffolding after "
            "apply_penum_normalization.py for [ED] commits."
        ),
    )
    args = parser.parse_args()

    model = Model.load(args.uml_path)
    written, skipped = 0, 0
    workbook = None
    penum_norm = load_penum_norm_csv(args.penum_norm_csv)

    def get_workbook():
        nonlocal workbook
        if workbook is None:
            workbook = load_property_workbook(args.properties_xlsx)
        return workbook

    for nm in args.names:
        element = model.by_name(nm)
        if element is None:
            workbook_content = get_workbook()
            if nm.startswith("Qto_") and nm in workbook_content.rows_by_set:
                kind = "qto"
            elif nm.startswith("Pset_") and nm in workbook_content.rows_by_set:
                kind = "pset"
            elif nm.startswith("PEnum_") and nm in workbook_content.penum_items:
                kind = "penum"
            else:
                print(f"[skip] {nm}: not found in {args.uml_path}", file=sys.stderr)
                skipped += 1
                continue
        elif isinstance(element, DataType):
            kind = "type"
        elif isinstance(element, Enumeration):
            kind = "penum" if nm.startswith("PEnum_") else "enum"
        elif isinstance(element, Class):
            kind = "qto" if nm.startswith("Qto_") else "pset" if nm.startswith("Pset_") else "entity"
        elif isinstance(element, Constraint) and element.language == "EXPRESS_FUNCTION":
            kind = "function"
        else:
            print(f"[skip] {nm}: unsupported XMI object {type(element).__name__}", file=sys.stderr)
            skipped += 1
            continue

        schema_name = _owning_schema_name(element) if element is not None else None

        if kind == "type":
            target = find_target_path(args.uml_path, args.docs_root, nm, kind, schema_name=schema_name)
            short_def = (
                load_short_def(args.ifc_spec_dir, nm, kind)
                or _extract_existing_short_def(target)
            )
            content = scaffold_datatype(nm, short_def=short_def)
            extra_outputs = []
        elif kind == "function":
            target = find_target_path(args.uml_path, args.docs_root, nm, kind, schema_name=schema_name)
            short_def = (
                load_short_def(args.ifc_spec_dir, nm, kind)
                or _extract_existing_short_def(target)
            )
            content = scaffold_function(nm, short_def=short_def)
            extra_outputs = []
        elif kind in {"enum", "penum"}:
            target = find_target_path(args.uml_path, args.docs_root, nm, kind, schema_name=schema_name)
            if kind == "penum":
                items = (
                    penum_norm.get(nm)
                    or get_workbook().penum_items.get(nm, [])
                )
                if items:
                    content = scaffold_penum(nm, items)
                else:
                    content = scaffold_penum(
                        nm,
                        [
                            PEnumWorkbookItem(nm, lit, "")
                            for lit in _model_enum_literals(element)
                        ],
                    )
            else:
                short_def = (
                    load_short_def(args.ifc_spec_dir, nm, kind)
                    or _extract_existing_short_def(target)
                )
                literal_docs = load_enum_literal_docs(args.ifc_spec_dir, nm)
                content = scaffold_enum(
                    nm, _model_enum_literals(element),
                    short_def=short_def, literal_docs=literal_docs,
                )
            extra_outputs = []
        elif kind in {"entity", "pset", "qto"}:
            target = find_target_path(args.uml_path, args.docs_root, nm, kind, schema_name=schema_name)
            short_def = (
                load_short_def(args.ifc_spec_dir, nm, kind)
                or _extract_existing_short_def(target)
            )
            extra_outputs = []
            if kind in {"pset", "qto"}:
                rows = get_workbook().rows_by_set.get(nm, [])
                comments = [
                    (row.name, row.local_definition)
                    for row in rows
                    if row.local_definition
                ]
                content = scaffold_property_or_quantity_set(
                    nm, short_def=short_def, comments=comments,
                )
                seen_prop_defs = {}
                for row in rows:
                    if row.name in seen_prop_defs and seen_prop_defs[row.name] != row.definition:
                        print(
                            f"[warn] {nm}: duplicate workbook definition for {row.name}; "
                            f"using first value",
                            file=sys.stderr,
                        )
                        continue
                    seen_prop_defs[row.name] = row.definition
                    prop_target = normalized_property_path(args.docs_root, row.name)
                    extra_outputs.append(
                        (prop_target, scaffold_normalized_property(row.name, row.definition))
                    )
            else:
                attr_docs, rule_docs = load_doc_entity(args.ifc_spec_dir, nm)
                if args.ifc_spec_dir and not (short_def or attr_docs or rule_docs):
                    print(
                        f"[info] {nm}: no prose found under --ifc-spec-dir; "
                        f"falling back to convention-based synthesis",
                        file=sys.stderr,
                    )
                content = scaffold_entity(
                    nm, None, _model_attrs(element), _model_rules(element),
                    attr_docs=attr_docs, rule_docs=rule_docs,
                    short_def=short_def,
                )
        else:
            print(f"[skip] {nm}: unsupported kind {kind}", file=sys.stderr)
            skipped += 1
            continue

        if target.exists() and not args.force:
            print(f"[skip] {target} exists; use --force to overwrite", file=sys.stderr)
            skipped += 1
            continue

        if args.dry_run:
            print(f"--- {target} ---")
            print(content)
            for extra_target, extra_content in extra_outputs:
                if extra_target.exists() and not args.force:
                    print(f"[skip] {extra_target} exists; use --force to overwrite", file=sys.stderr)
                    continue
                print(f"--- {extra_target} ---")
                print(extra_content)
        else:
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(content, encoding="utf-8")
            print(f"[wrote] {target}")
            written += 1
            for extra_target, extra_content in extra_outputs:
                if extra_target.exists() and not args.force:
                    print(f"[skip] {extra_target} exists; use --force to overwrite", file=sys.stderr)
                    skipped += 1
                    continue
                extra_target.parent.mkdir(parents=True, exist_ok=True)
                extra_target.write_text(extra_content, encoding="utf-8")
                print(f"[wrote] {extra_target}")
                written += 1

    print(f"\n{written} written, {skipped} skipped.", file=sys.stderr)


if __name__ == "__main__":
    main()
